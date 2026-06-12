from openpyxl import load_workbook
from openpyxl.styles import Alignment

wb = load_workbook('/Users/wendytan/iac-i2c-v2-build-hackathon/i2cv2-resource-coverage.xlsx')
ws = wb['Resource Coverage']

max_row = ws.max_row
start_row = 2
current_val = ws.cell(row=2, column=1).value

for row in range(3, max_row + 2):
    val = ws.cell(row=row, column=1).value if row <= max_row else None
    if val != current_val:
        if row - 1 > start_row:
            ws.merge_cells(start_row=start_row, start_column=1, end_row=row - 1, end_column=1)
            ws.cell(row=start_row, column=1).alignment = Alignment(vertical='center', horizontal='center')
        start_row = row
        current_val = val

wb.save('/Users/wendytan/iac-i2c-v2-build-hackathon/i2cv2-resource-coverage.xlsx')
print('Done - column A cells merged by service name.')
